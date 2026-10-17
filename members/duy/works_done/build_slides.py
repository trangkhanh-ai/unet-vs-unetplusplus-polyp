"""Create presentation slides from the project's completed locked experiments.

Uses the bundled document Python runtime with python-pptx, not the torch venv.
Run directly: python members/duy/works_done/build_slides.py
--validate-only reads evidence and PNG metadata without creating an artifact.
The authoring caller must perform artifact marking and visual QA separately.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

if __package__:
    from .build_report import Evidence, PROJECT_ROOT, load_evidence, png_size, score, integer, training_precision, author_date
else:
    # Direct execution avoids the model package's __init__ and its torch import.
    from build_report import Evidence, PROJECT_ROOT, load_evidence, png_size, score, integer, training_precision, author_date


WIDTH, HEIGHT = 13.333333, 7.5
NAVY, BLACK, ACCENT, GREY, WHITE = "172F44", "111111", "17855C", "D9D9D9", "FFFFFF"
FONT = "Calibri"


def required_images(root: Path, comparison: Path) -> dict[str, Path]:
    report = root / "results/report_figures"
    paths = {
        "curves": report / "training_curves.png",
        "paired": report / "paired_examples.png",
    }
    if not paths["curves"].is_file():
        paths["curves"] = comparison.parent / "training_curves.png"
    if not paths["paired"].is_file():
        paths["paired"] = comparison.parent / "paired_examples.png"
    for path in paths.values():
        if not path.is_file():
            raise FileNotFoundError(f"Thiếu hình trình bày: {path}")
        width, height = png_size(path)
        if width < 1 or height < 1:
            raise ValueError(f"Kích thước PNG không hợp lệ: {path}")
    return paths


def add_text(slide, text: str, x: float, y: float, w: float, h: float,
             *, size: float = 22, bold: bool = False, color: str = BLACK,
             align: str = "left", gap: float = 9):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = alignment
        paragraph.space_after = Pt(gap)
        paragraph.space_before = Pt(0)
        paragraph.line_spacing = 1.08
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
    return box


def add_box(slide, text: str, x: float, y: float, w: float, h: float,
            *, size: float = 22, color: str = NAVY, border: str = GREY):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    shape.line.color.rgb = RGBColor.from_string(border)
    shape.line.width = Pt(1.2)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.12)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, color: str = GREY):
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = RGBColor.from_string(color)
    line.line.width = Pt(2)
    return line


def add_picture_fit(slide, path: Path, x: float, y: float, w: float, h: float):
    from pptx.util import Inches
    width, height = png_size(path)
    scale = min(w / width, h / height)
    display_w, display_h = width * scale, height * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - display_w) / 2),
                                    Inches(y + (h - display_h) / 2),
                                    width=Inches(display_w), height=Inches(display_h))


def add_unet_diagram(slide) -> None:
    """Draw editable native nodes at presentation size, rather than a tiny PNG."""
    left, right, box_width, box_height = 0.8, 5.3, 2.7, 0.76
    rows = [1.55 + level * 0.99 for level in range(5)]
    # Draw connections first, so their ends remain behind the boxes.
    for level in range(4):
        y = rows[level]
        add_line(slide, left + box_width / 2, y + box_height,
                 left + box_width / 2, rows[level + 1], GREY)
        add_line(slide, left + box_width, y + box_height / 2,
                 right, y + box_height / 2, ACCENT)
        if level < 3:
            add_line(slide, right + box_width / 2, rows[level + 1],
                     right + box_width / 2, y + box_height, NAVY)
    add_line(slide, left + box_width, rows[4] + box_height / 2,
             right + box_width / 2, rows[3] + box_height, NAVY)
    for level, (channels, resolution) in enumerate(zip((32, 64, 128, 256, 512), (256, 128, 64, 32, 16))):
        add_box(slide, f"{channels} ch | {resolution} × {resolution}",
                left, rows[level], box_width, box_height, size=21)
        if level < 4:
            add_box(slide, f"{channels} ch | {resolution} × {resolution}",
                    right, rows[level], box_width, box_height, size=21)
    add_text(slide, "Encoder ↓", 0.8, 1.2, 2.7, 0.35, size=20, bold=True, color=NAVY, gap=0)
    add_text(slide, "Decoder ↑", 5.3, 1.2, 2.7, 0.35, size=20, bold=True, color=NAVY, gap=0)
    add_text(slide, "Đường xanh: concat skip trực tiếp", 0.8, 6.7, 7.2, 0.4, size=20, color=ACCENT, gap=0)


def add_unetplusplus_diagram(slide) -> None:
    """Show all 15 nodes; omit older dense edges explicitly to keep it readable."""
    origin_x, origin_y = 0.65, 1.6
    dx, dy, box_width, box_height = 1.65, 0.89, 1.38, 0.67
    position = {(i, j): (origin_x + j * dx, origin_y + i * dy)
                for i in range(5) for j in range(5 - i)}
    for i in range(4):
        x, y = position[i, 0]
        lower_x, lower_y = position[i + 1, 0]
        add_line(slide, x + box_width / 2, y + box_height,
                 lower_x + box_width / 2, lower_y, GREY)
    for i in range(4):
        for j in range(1, 5 - i):
            x, y = position[i, j]
            previous_x, previous_y = position[i, j - 1]
            lower_x, lower_y = position[i + 1, j - 1]
            add_line(slide, previous_x + box_width, previous_y + box_height / 2,
                     x, y + box_height / 2, ACCENT)
            add_line(slide, lower_x + box_width / 2, lower_y,
                     x + box_width / 2, y + box_height, NAVY)
    for (i, j), (x, y) in position.items():
        add_box(slide, f"X({i},{j})", x, y, box_width, box_height,
                size=21, border=ACCENT if (i, j) == (0, 4) else GREY)
    add_text(slide, "Cột j=0 là encoder; i tăng khi resolution giảm", 0.65, 6.65, 8.4, 0.36,
             size=20, color=NAVY, gap=0)
    add_text(slide, "Sơ đồ giản lược: chỉ vẽ cạnh ngang liền kề.\nDense skip còn nhận mọi node cũ cùng level.",
             0.65, 6.0, 8.4, 0.6, size=20, gap=0)


def new_slide(presentation, number: int, title: str, notes: str):
    from pptx.dml.color import RGBColor
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    add_text(slide, title, 0.55, 0.3, 12.2, 0.85, size=30, bold=True, color=NAVY)
    add_line(slide, 0.58, 1.15, 12.75, 1.15)
    add_text(slide, "Đồ án Deep Learning  •  Kvasir-SEG  •  " + author_date(), 0.58, 7.13, 10.7, 0.22, size=12, color=NAVY, gap=0)
    add_text(slide, str(number), 12.0, 7.08, 0.7, 0.35, size=14, color=NAVY, align="right", gap=0)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_results_table(slide, evidence: Evidence):
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    headers = ["Model", "Dice", "IoU", "Tham số", "Epoch best"]
    names = ["A  U-Net", "B  U-Net++", "C  U-Net++ DS"]
    rows = [[name, score(row["dice"]), score(row["iou"]), integer(row["parameters"]), str(row["selected_epoch"])]
            for name, row in zip(names, evidence.rows)]
    table_shape = slide.shapes.add_table(4, 5, Inches(0.75), Inches(1.7), Inches(11.85), Inches(3.0))
    table = table_shape.table
    for column, width in zip(table.columns, (3.35, 1.55, 1.55, 2.55, 2.85)):
        column.width = Inches(width)
    table.first_row = False
    for index, values in enumerate([headers, *rows]):
        for col, value in enumerate(values):
            cell = table.cell(index, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("EDEDED" if index == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame = cell.text_frame
            frame.clear()
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT if col == 0 else PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(21)
            run.font.bold = index == 0
            run.font.color.rgb = RGBColor.from_string(NAVY if index == 0 else BLACK)


def create_slides(e: Evidence, images: dict[str, Path], output: Path) -> int:
    """Only this function authors a PPTX. The caller handles the marker and QA."""
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(WIDTH), Inches(HEIGHT)
    presentation.core_properties.title = "Phân đoạn polyp bằng U-Net và U-Net++ trên Kvasir-SEG"
    presentation.core_properties.subject = "Bảo vệ đồ án môn Deep Learning"
    presentation.core_properties.author = "Nhóm đồ án Deep Learning"
    a, b, c = e.rows
    cfg = e.runs[0]["config"]
    precision_label, precision_description = training_precision(cfg)
    precision_history = (
        " Run thử C gặp overflow forward FP16 tại batch 28, epoch 30. Trọng số và forward FP32 vẫn hữu hạn. "
        "Sau chẩn đoán train, cả A/B/C chạy lại từ đầu với BF16 trước khi khóa hoặc đọc test; không tune theo điểm test."
        if cfg.get("amp") and cfg.get("amp_dtype") == "bfloat16" else ""
    )
    delta_b = float(b["dice"]) - float(a["dice"])
    delta_c = float(c["dice"]) - float(b["dice"])
    best = max(e.rows, key=lambda row: float(row["dice"]))

    slide = new_slide(presentation, 1, "Phân đoạn polyp bằng U-Net và U-Net++",
        f"Mở đầu: đề tài yêu cầu nghiên cứu U-Net++, nested dense skip và so sánh với U-Net trên Kvasir-SEG. "
        f"Ba run đã hoàn thành; kết quả đọc từ artifact. Cấu hình có Dice cao nhất là {best['experiment']} "
        f"với {score(best['dice'])}, chỉ trong split và seed này. Không hứa trước B hay C phải thắng.")
    add_text(slide, "Kvasir-SEG  |  Đồ án môn Deep Learning", 0.85, 1.6, 11.65, 0.55, size=26, color=NAVY)
    add_text(slide, "Ảnh nội soi RGB → mask vùng polyp", 0.85, 2.4, 11.65, 0.75, size=32, bold=True)
    add_text(slide, "A  U-Net baseline\nB  U-Net++ với nested dense skip\nC  U-Net++ thêm deep supervision", 0.9, 3.55, 11.4, 2.0, size=26, gap=14)
    add_text(slide, "Cùng 150 ảnh test  •  seed 42  •  checkpoint chọn bằng validation", 0.9, 6.08, 11.4, 0.6, size=22, color=ACCENT)

    slide = new_slide(presentation, 2, "Dữ liệu và phạm vi bài toán",
        "Output là semantic segmentation nhị phân, không phải chẩn đoán hay bounding box. Dataset có 1000 cặp. "
        "Audit kiểm tra đọc được, kích thước và pixel trùng hoàn toàn; chưa xác minh ảnh gần trùng hoặc patient/video independence. "
        "Train/val/test giữ cố định theo manifest SHA256. Mask JPEG được threshold >127 trước nearest resize.")
    for x, label, count in ((0.8, "Train", 700), (4.65, "Validation", 150), (8.5, "Test", 150)):
        add_box(slide, f"{label}\n{count} ảnh", x, 1.65, 3.35, 1.15, size=26)
    add_text(slide, "RGB 256 × 256, chia 255\nMask 0/1, resize nearest\nAugmentation chỉ trên train", 0.85, 3.45, 5.85, 2.7, size=24, gap=17)
    add_text(slide, "Train cập nhật trọng số\nValidation chọn best\nTest sau khi khóa protocol", 7.0, 3.45, 5.45, 2.7, size=24, gap=17)
    add_text(slide, "Không dùng cột kfold_5 trong thực nghiệm holdout này", 0.85, 6.28, 11.65, 0.45, size=20, color=NAVY)

    slide = new_slide(presentation, 3, "U-Net nối trực tiếp encoder và decoder",
        "Đi theo một ảnh từ 256 xuống 16 rồi trở lại 256. Mỗi block có hai Conv 3×3 và ReLU. "
        "Bốn skip concat giữ chi tiết vị trí. ConvTranspose2d đưa feature về resolution của skip. "
        "Topology U-Net được giữ; same padding/base 32 là lựa chọn triển khai, không tái lập nguyên xi paper 2015.")
    add_unet_diagram(slide)
    add_text(slide, "256 → 16 → 256\n\n4 skip trực tiếp\n\nHead 1 channel\ntrả logits", 9.0, 1.8, 3.7, 4.6, size=24, gap=10)
    add_text(slide, "Same padding  •  base 32\nkhông BatchNorm", 8.95, 6.18, 3.8, 0.7, size=20, color=ACCENT)

    slide = new_slide(presentation, 4, "U-Net++ tinh chỉnh feature trên đường skip",
        "Chỉ cột j=0 là encoder. i là độ sâu, j là cột tinh chỉnh. Mỗi node nhận tất cả node trước ở cùng level "
        "và feature upsample từ dưới. Động cơ paper là giảm semantic gap, nhưng thí nghiệm này không đo trực tiếp semantic gap. "
        "B lấy head X(0,4); C giám sát bốn head, suy luận vẫn dùng head cuối. Đồ án không triển khai average mode hoặc pruning.")
    add_unetplusplus_diagram(slide)
    add_text(slide, "Nested\nĐường skip lồng nhau\n\nDense\nGiữ node cùng level\n\nB/C dùng X(0,4)", 9.1, 1.8, 3.65, 4.9, size=22, gap=10)

    slide = new_slide(presentation, 5, "Ví dụ đọc node X(0,2)",
        "Đây là ví dụ cần tự giải thích. X(0,2) nối X(0,0) và X(0,1) với Up(X(1,1)). "
        "Ở base 32, Up chiếu 64 thành 32 channel, vì vậy concat nhận 96 channel rồi DoubleConv trả 32 channel. "
        "torch.cat dim=1 là channel, không tăng H/W. Node mới có trọng số riêng. X(0,4) tương tự nhận 160 channel.")
    inputs = [("X(0,0)\n32 channel", 1.6), ("X(0,1)\n32 channel", 3.05), ("Up(X(1,1))\n64 → 32 channel", 4.5)]
    for text, y in inputs:
        add_box(slide, text, 0.85, y, 3.25, 1.05, size=24)
        add_line(slide, 4.12, y + 0.53, 5.0, 3.45, ACCENT)
    add_box(slide, "Concat dim=1\n96 channel", 5.05, 2.7, 3.0, 1.5, size=26, border=ACCENT)
    add_text(slide, "→", 8.2, 3.08, 0.6, 0.6, size=32, bold=True, color=ACCENT)
    add_box(slide, "DoubleConv\nX(0,2): 32 channel", 9.0, 2.7, 3.35, 1.5, size=24)
    add_text(slide, "H/W vẫn 256 × 256  •  Dense skip giữ toàn bộ node cũ", 0.85, 6.05, 11.7, 0.75, size=24, color=NAVY)

    slide = new_slide(presentation, 6, "Loss và cách model học",
        "Model trả logits. BCEWithLogitsLoss kết hợp sigmoid ổn định; không sigmoid trước hàm loss này. "
        "Soft Dice dùng xác suất và epsilon 1e-6, tính từng ảnh rồi mean. Adam cập nhật trọng số bằng gradient của autograd. "
        "C lấy trung bình bốn loss; val/test chỉ dùng head cuối. " + precision_description + " Validation/test dùng FP32.")
    add_text(slide, "Loss = 0,5 BCEWithLogits + 0,5 soft Dice loss", 0.8, 1.7, 11.75, 0.7, size=28, bold=True, color=NAVY)
    for index, word in enumerate(("Batch", "Forward", "Loss", "Backward", "Adam")):
        x = 0.8 + index * 2.46
        add_box(slide, word, x, 2.95, 2.1, 1.0, size=24)
        if index < 4:
            add_text(slide, "→", x + 2.12, 3.13, 0.35, 0.55, size=28, color=ACCENT)
    add_text(slide, "Sau epoch: validation → chọn best → lưu checkpoint", 0.8, 4.6, 11.75, 0.7, size=25)
    add_text(slide, "C: trung bình loss 4 head; inference dùng head cuối", 0.8, 5.7, 11.75, 0.65, size=24, color=ACCENT)

    slide = new_slide(presentation, 7, "Giữ chung protocol để so sánh",
        "Giữ cùng batch và điều kiện A/B/C. B/C có cùng trunk và head cuối khởi tạo theo seed 42. "
        "Sampler và augmentation reset seed riêng theo epoch. Best lấy Dice validation cao nhất, hòa thì loss thấp nhất. "
        "Protocol lock ghi SHA của checkpoint/config/manifest trước test. Timestamp và file lock chỉ là dấu vết trong repo." + precision_history)
    add_text(slide, f"Base {cfg['model_kwargs']['base_channels']}  •  RGB 256 × 256\nAdam lr {cfg['learning_rate']}  •  batch {cfg['batch_size']}\nTối đa {cfg['max_epochs']} epoch  •  patience {cfg['early_stopping_patience']}\nSeed {cfg['seed']}  •  cùng augmentation", 0.8, 1.65, 6.0, 4.55, size=25, gap=22)
    add_text(slide, f"Train: {precision_label}\nValidation/test: FP32\nKhông TF32\nThreshold 0,5 cố định", 7.0, 1.65, 5.5, 4.55, size=25, gap=22)
    add_text(slide, "Không pretrained  •  Không BatchNorm  •  Best được khóa trước test", 0.8, 6.3, 11.9, 0.45, size=20, color=ACCENT)

    slide = new_slide(presentation, 8, "Kết quả trên cùng 150 ảnh test",
        f"Số liệu thật, metric lấy mean từng ảnh ở resolution 256. Dice B−A={delta_b:+.4f}; C−B={delta_c:+.4f}. "
        "Nếu B thấp hơn A, báo đúng kết quả và nghiên cứu giới hạn. Không điều chỉnh cấu hình sau khi xem test. "
        "A/B khác số tham số nên chưa cô lập riêng topology. Một seed chưa đủ để báo std qua nhiều lần hoặc khẳng định ý nghĩa thống kê.")
    add_results_table(slide, e)
    add_text(slide, f"ΔDice B−A ={delta_b:+.4f}    |    ΔDice C−B ={delta_c:+.4f}", 0.85, 5.25, 11.65, 0.6, size=26, bold=True, color=NAVY)
    add_text(slide, "Mean theo ảnh  •  threshold 0,5  •  head cuối  •  một seed", 0.85, 6.15, 11.65, 0.5, size=22, color=ACCENT)

    slide = new_slide(presentation, 9, "Đọc đường học và chọn checkpoint",
        "Đường validation Dice chọn best epoch trước khi đọc test. Train loss C là mean 4 head còn val loss là head cuối, nên hai đường khác objective. "
        "Train có augmentation còn val không có. Cần đọc validation Dice và mask cùng với khoảng cách loss. Last checkpoint phục vụ resume; test dùng best. "
        "training_and_validation_seconds cộng wall time qua epoch, có thể gồm thời gian tạm dừng trong epoch, không đo GPU compute. "
        "Run B có gián đoạn do máy ngủ, tiến trình tiếp tục với trọng số và trạng thái còn trong bộ nhớ; không đổi thuật toán hay cấu hình. "
        "Không kết luận tốc độ train từ wall time. Timing forward test dùng cùng GPU và FP32 theo scope đã ghi, nhưng chỉ đo một lượt và có thể dao động theo trạng thái GPU. "
        "Lượt này C nhanh hơn B theo số đo; B/C có cùng graph chính và C tính thêm ba head phụ. Không suy ra deep supervision giảm thời gian suy luận hoặc model đã pruning.")
    add_picture_fit(slide, images["curves"], 0.55, 1.35, 6.45, 5.5)
    add_text(slide, "Best theo validation Dice\n\nTest dùng best.pt\n\nC: train mean 4 head\nC: val head cuối", 7.3, 1.9, 5.35, 4.6, size=24, gap=12)
    add_text(slide, "Số epoch thực tế có thể khác do cùng quy tắc dừng sớm", 7.25, 6.1, 5.3, 0.7, size=20, color=ACCENT)

    is_fixed_ids = images["paired"].parent.name == "report_figures"
    selection = "Ba test ID đầu theo thứ tự sắp xếp; không chọn theo điểm" if is_fixed_ids else "Các ca cực trị theo delta Dice và mean Dice; chọn sau đánh giá"
    slide = new_slide(presentation, 10, "Xem mask và phân tích lỗi",
        selection + ". Đọc mỗi hàng cùng ảnh và GT, so vùng bỏ sót, nhầm nền và sai biên. "
        "Không dùng các ảnh này tune lại rồi gọi là test độc lập. Bảng paired_per_image giữ toàn bộ 150 ID; mask PNG đo ở 256. "
        "Cần xem toàn bộ bảng, thêm các ca cực trị và nhóm diện tích trước khi kết luận.")
    add_picture_fit(slide, images["paired"], 0.65, 1.35, 12.0, 5.0)
    add_text(slide, selection, 0.85, 6.4, 11.65, 0.45, size=20, color=NAVY)

    slide = new_slide(presentation, 11, "Kết luận chỉ áp dụng trong điều kiện đã ghi",
        "Nhóm small/medium/large là tỷ lệ diện tích mask trên ảnh, không phân loại kích thước bằng mm. "
        "Audit không phát hiện ảnh trùng hoàn toàn theo pixel giải mã; ảnh gần trùng và độc lập patient/video chưa xác minh. "
        "Mask JPEG và resize có thể thay đổi biên. Một Kvasir holdout chưa đủ cho kết luận lâm sàng. Nghiên cứu tiếp theo nên dùng nhiều seed và dữ liệu ngoài.")
    add_text(slide, "• Seed 42, một split holdout\n• A/B khác số tham số\n• Chưa xác minh patient/video và near duplicates\n• Nhóm diện tích ảnh không phải kích thước lâm sàng\n• Chưa đánh giá trên dataset hoặc cơ sở y tế khác", 0.85, 1.7, 11.65, 4.4, size=25, gap=17)
    add_text(slide, "Hướng tiếp theo: nhiều seed và dữ liệu ngoài", 0.85, 6.2, 11.65, 0.65, size=25, bold=True, color=ACCENT)

    slide = new_slide(presentation, 12, "Demo dùng đúng checkpoint đã chọn",
        "Duy mở demo theo README, tải ảnh hoặc chọn ví dụ. Chọn A/B/C dùng đúng best checkpoint đã khóa. "
        "Không có GT thì demo không tính Dice/IoU cho ảnh upload. Ground truth phải tương ứng thì mới đánh giá. "
        "Model trả logits; demo sigmoid, resize xác suất bilinear rồi threshold trên ảnh gốc. C chọn head cuối. Timing forward của báo cáo không đo độ trễ đầu cuối của web demo.")
    add_text(slide, "1  Mở demo và chọn model A/B/C\n2  Tải ảnh nội soi hoặc chọn ảnh minh họa\n3  Xem mask và overlay ở ảnh gốc\n4  Giải thích vùng đúng và vùng bị bỏ sót", 0.85, 1.75, 11.65, 3.6, size=27, gap=18)
    add_text(slide, "Ảnh upload không có ground truth → không có Dice/IoU", 0.85, 5.65, 11.65, 0.9, size=27, bold=True, color=ACCENT)

    slide = new_slide(presentation, 13, "Năm phần code và phần cần tự giải thích",
        "Đây là phân công module; mỗi bạn cần tự chạy và giải thích phần mình. "
        "Duy cần hiểu luồng dữ liệu → model → loss → best → test và giới hạn so sánh. "
        "Các test tự động kiểm tra giao diện và logic; hiệu quả phân đoạn thực tế được đo riêng bằng các metrics test.")
    add_text(slide, "Duy       U-Net, tích hợp, protocol và báo cáo\nKhánh   Dataset, mask, split và augmentation\nPhúc     Nested dense skip và deep supervision\nHuy       Loss, train, validation và checkpoint\nTrí         Metrics, paired analysis và demo", 0.85, 1.7, 11.65, 3.75, size=26, gap=17)
    add_text(slide, f"{e.tests['tests_run']}/{e.tests['tests_run']} kiểm tra đạt  •  mỗi bạn tự chạy và giải thích module", 0.85, 6.08, 11.65, 0.65, size=24, color=ACCENT)

    slide = new_slide(presentation, 14, "Tóm tắt kết quả và nguồn học",
        f"Yêu cầu A/B đã thực hiện; C nghiên cứu mở rộng. Trong run này, best test Dice={score(best['dice'])} ở {best['experiment']}. "
        "Kết quả chỉ áp dụng cho cấu hình và split đã ghi. Đọc U-Net 2015, U-Net++ 2018, Kvasir-SEG 2019 và tài liệu PyTorch. "
        "Trước bảo vệ, Duy cần tự tính Dice ví dụ, giải thích X(0,2), đọc best epoch và demo ảnh không có GT."
    )
    add_text(slide, "Đã triển khai A/B và mở rộng deep supervision\nKết quả có checkpoint, CSV, mask và provenance\nCần đọc giới hạn trước khi kết luận model tốt hơn", 0.85, 1.6, 11.65, 2.4, size=26, gap=16)
    add_text(slide, "U-Net     arxiv.org/abs/1505.04597\nU-Net++  arxiv.org/abs/1807.10165\nKvasir-SEG  arxiv.org/abs/1911.07069\nPyTorch    docs.pytorch.org", 0.85, 4.25, 11.65, 2.1, size=22, color=NAVY, gap=8)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return len(presentation.slides)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--comparison", type=Path, default=Path("results/comparison/comparison.json"))
    parser.add_argument("--output", type=Path, default=Path("reports/SLIDES_UNET_UNETPLUSPLUS_KVASIR_SEG.pptx"))
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    root = args.project_root.resolve()
    comparison = args.comparison if args.comparison.is_absolute() else root / args.comparison
    evidence = load_evidence(root, comparison)
    images = required_images(root, comparison)
    if args.validate_only:
        print(json.dumps({"status": "slides_evidence_validated", "planned_slides": 14,
                          "figures": {name: str(path) for name, path in images.items()}}, ensure_ascii=False, indent=2))
        return
    output = args.output if args.output.is_absolute() else root / args.output
    count = create_slides(evidence, images, output)
    print(json.dumps({"status": "slides_created_pending_visual_qa", "slides": count, "output": str(output)},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
