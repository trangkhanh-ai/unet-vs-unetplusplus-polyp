"""Build the course report from completed experiments, without importing torch.

Run this authoring script with the bundled document runtime, after comparison:
    python members/duy/works_done/build_report.py

Use this direct script path: running through the works_done package also loads
the PyTorch model package, which the document runtime deliberately does not need.

--validate-only checks inputs without creating a Word or Markdown artifact.
No scores are invented. Missing or unfinished evidence stops report creation.
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
import hashlib
import json
import math
from pathlib import Path
import struct
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[3]
RUN_NAMES = ("A_unet_seed42", "B_unetplusplus_seed42", "C_unetplusplus_ds_seed42")
MODEL_NAMES = ("U-Net", "U-Net++", "U-Net++ có deep supervision")
TITLE = "Phân đoạn polyp bằng UNet và UNetPlusPlus trên Kvasir SEG"
REFERENCES = (
    ("1", "Ronneberger O, Fischer P, Brox T. U-Net: Convolutional Networks for Biomedical Image Segmentation. 2015.", "https://arxiv.org/abs/1505.04597"),
    ("2", "Zhou Z, Siddiquee M M R, Tajbakhsh N, Liang J. UNet++: A Nested U-Net Architecture for Medical Image Segmentation. 2018.", "https://arxiv.org/abs/1807.10165"),
    ("3", "Jha D và cộng sự. Kvasir-SEG: A Segmented Polyp Dataset. 2019.", "https://arxiv.org/abs/1911.07069"),
    ("4", "PyTorch. BCEWithLogitsLoss. API tham khảo cho phiên bản 2.11.", "https://docs.pytorch.org/docs/2.11/generated/torch.nn.BCEWithLogitsLoss.html"),
    ("5", "PyTorch. Learn the Basics. Tensor, Dataset, autograd và optimization.", "https://docs.pytorch.org/tutorials/beginner/basics/intro.html"),
    ("6", "PyTorch. Saving and Loading Models. State dict và checkpoint.", "https://docs.pytorch.org/tutorials/beginner/saving_loading_models.html"),
    ("7", "PyTorch. Reproducibility. Giới hạn của seed và deterministic algorithms.", "https://docs.pytorch.org/docs/2.11/notes/randomness.html"),
)


def author_date() -> str:
    """Show the artifact creation date in the user's UTC+7 timezone."""
    return datetime.now(timezone(timedelta(hours=7))).strftime("%d/%m/%Y")


def read_json(path: Path) -> dict:
    if not path.is_file():
        raise FileNotFoundError(f"Thiếu bằng chứng cần thiết: {path}")
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"JSON phải là object: {path}")
    return value


def read_csv(path: Path) -> list[dict]:
    if not path.is_file():
        raise FileNotFoundError(f"Thiếu bảng cần thiết: {path}")
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    if not rows:
        raise ValueError(f"Bảng rỗng: {path}")
    return rows


def finite(value: Any) -> float:
    number = float(value)
    if not math.isfinite(number):
        raise ValueError(f"Giá trị không hữu hạn: {value}")
    return number


def score(value: Any) -> str:
    return f"{finite(value):.4f}"


def integer(value: Any) -> str:
    return f"{int(value):,}".replace(",", ".")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def training_precision(config: dict) -> tuple[str, str]:
    """Read the precision from the finished run, so prose matches its config."""
    if not config.get("amp", False):
        return "FP32", "Train forward/backward dùng FP32, không dùng AMP."
    dtype = config.get("amp_dtype", "float16")
    if dtype == "bfloat16":
        return "AMP BF16", "Train forward dùng AMP bfloat16 (BF16), không dùng GradScaler cho BF16."
    if dtype == "float16":
        return "AMP FP16", "Train forward dùng AMP float16 (FP16) và GradScaler."
    raise ValueError(f"amp_dtype không được hỗ trợ trong báo cáo: {dtype}")


@dataclass
class Evidence:
    root: Path
    comparison_path: Path
    audit: dict
    tests: dict
    comparison: dict
    runs: list[dict]
    histories: list[list[dict]]
    groups: list[dict]
    paired: list[dict]
    lock: dict

    @property
    def rows(self) -> list[dict]:
        return self.comparison["comparison"]

    @property
    def labels(self) -> list[str]:
        return self.comparison["labels"]


def load_evidence(root: Path, comparison_path: Path) -> Evidence:
    """Validate finished runs and matching evidence before authoring anything."""
    audit = read_json(root / "results/data_audit/audit.json")
    tests = read_json(root / "results/verification/test_summary.json")
    comparison = read_json(comparison_path)
    lock = read_json(root / "configs/protocol_lock.json")
    runs = [read_json(root / "runs" / name / "summary.json") for name in RUN_NAMES]
    histories = [read_csv(root / "runs" / name / "history.csv") for name in RUN_NAMES]
    groups = read_csv(comparison_path.parent / "comparison_by_size.csv")
    paired = read_csv(comparison_path.parent / "paired_per_image.csv")
    comparison_csv = read_csv(comparison_path.parent / "comparison.csv")
    if audit.get("status") != "passed" or audit.get("issues"):
        raise ValueError("Data audit chưa đạt")
    if tests.get("status") != "passed" or tests.get("failure_count") or tests.get("error_count"):
        raise ValueError("Kiểm tra dự án chưa đạt")
    if comparison.get("status") != "compared" or comparison.get("split") != "test":
        raise ValueError("Báo cáo chính thức cần comparison đã hoàn thành trên test")
    labels, rows = comparison.get("labels", []), comparison.get("comparison", [])
    if len(labels) != 3 or len(set(labels)) != 3 or len(rows) != 3:
        raise ValueError("Cần ba thí nghiệm A/B/C, mỗi thí nghiệm có label riêng")
    if len(comparison_csv) != 3 or len(paired) != 150 or comparison.get("image_count") != 150:
        raise ValueError("Cần đúng ba cấu hình và 150 ảnh test")
    if len({row["image_id"] for row in paired}) != len(paired):
        raise ValueError("Bảng paired có ID bị trùng")
    manifest_hash = comparison["manifest_sha256"]
    if audit["manifest_sha256"] != manifest_hash or lock["manifest_sha256"] != manifest_hash:
        raise ValueError("Audit, khóa và kết quả dùng manifest khác nhau")
    manifest = root / runs[0]["config"]["manifest"]
    if sha256(manifest) != manifest_hash:
        raise ValueError("Manifest hiện tại đã thay đổi")
    if audit["split_counts"] != {"train": 700, "val": 150, "test": 150}:
        raise ValueError("Báo cáo này yêu cầu holdout 700/150/150")
    if comparison["protocol_lock_sha256"] != sha256(root / "configs/protocol_lock.json"):
        raise ValueError("Kết quả không khớp SHA256 của protocol lock")
    if lock.get("status") != "locked_before_test" or labels != ["A", "B", "C"]:
        raise ValueError("Báo cáo yêu cầu khóa trước test và thứ tự A/B/C đã chốt")
    locked = {entry["label"]: entry for entry in lock.get("experiments", [])}
    if set(locked) != set(labels):
        raise ValueError("Protocol lock không có đủ ba thí nghiệm A/B/C")
    test_ids = {row["image_id"] for row in read_csv(manifest) if row["split"] == "test"}
    if {row["image_id"] for row in paired} != test_ids:
        raise ValueError("Bảng paired không chứa đúng test ID trong manifest")
    source_hashes = runs[0]["provenance"]["source_hashes"]
    for name, expected_hash in source_hashes.items():
        if sha256(root / name) != expected_hash:
            raise ValueError(f"Source thực thi đã thay đổi sau huấn luyện: {name}")
    for index, (row, run, history) in enumerate(zip(rows, runs, histories)):
        if run.get("status") not in {"completed", "early_stopped"}:
            raise ValueError(f"Run chưa kết thúc: {RUN_NAMES[index]}")
        if run["test_images_read"] != 0 or run["manifest_sha256"] != manifest_hash:
            raise ValueError("Trainer phải chỉ đọc train/val từ manifest đã chốt")
        if row["experiment"] != labels[index] or row["parameters"] != run["parameter_count"]:
            raise ValueError("Bảng comparison không khớp run")
        if row["selected_epoch"] != run["best"]["epoch"]:
            raise ValueError("Kết quả không dùng checkpoint chọn bằng validation")
        entry = locked[labels[index]]
        checkpoint = root / "runs" / RUN_NAMES[index] / "best.pt"
        if checkpoint.resolve() != (root / entry["checkpoint_path"]).resolve():
            raise ValueError("Protocol lock trỏ đến checkpoint khác run báo cáo")
        if sha256(checkpoint) != entry["checkpoint_sha256"] or entry["config"] != run["config"]:
            raise ValueError("Checkpoint hoặc cấu hình đã khác protocol lock")
        if run["provenance"]["source_hashes"] != source_hashes:
            raise ValueError("Các run dùng source thực thi khác nhau")
        expected_ds = index == 2
        if bool(row["deep_supervision"]) != expected_ds:
            raise ValueError("Thứ tự cấu hình phải là U-Net, U-Net++, U-Net++ deep supervision")
        expected_model = "unet" if index == 0 else "unetplusplus"
        if row["model"] != expected_model:
            raise ValueError("Thứ tự kiến trúc trong comparison không đúng")
        if int(history[-1]["epoch"]) != run["epochs_completed"]:
            raise ValueError("Số epoch của history và summary khác nhau")
        for name in ("dice", "iou", "precision", "recall"):
            value = finite(row[name])
            if not 0 <= value <= 1:
                raise ValueError(f"Metric ngoài [0,1]: {name}")
            values = [finite(image[f"{name}_{labels[index]}"]) for image in paired]
            if not math.isclose(sum(values) / len(values), value, abs_tol=1e-8, rel_tol=0):
                raise ValueError("Bảng paired và summary không khớp")
            if not math.isclose(finite(comparison_csv[index][name]), value, abs_tol=1e-10, rel_tol=0):
                raise ValueError("comparison.csv và comparison.json không khớp")
        own_groups = [group for group in groups if group["experiment"] == labels[index]]
        if sum(int(group["count"]) for group in own_groups) != len(paired):
            raise ValueError("Nhóm diện tích không phủ đủ test ID")
    return Evidence(root, comparison_path, audit, tests, comparison, runs, histories, groups, paired, lock)


# A small content model keeps Word and Markdown statements identical.
# Blocks contain human-readable paragraphs, concise tables and native equations.
Block = tuple[str, Any]


def p(text: str) -> Block:
    return "paragraph", text


def h(text: str) -> Block:
    return "subheading", text


def table(headers: list[str], rows: list[list[str]], widths: list[float], numeric: tuple[int, ...] = ()) -> Block:
    return "table", {"headers": headers, "rows": rows, "widths": widths, "numeric": numeric}


def figure(path: Path, caption: str, width: float = 6.5, max_height: float = 6.05) -> Block:
    if not path.is_file():
        raise FileNotFoundError(f"Thiếu hình báo cáo: {path}")
    return "figure", {"path": path, "caption": caption, "width": width, "max_height": max_height}


def compact_figure(evidence: Evidence, filename: str, fallback: str) -> Path:
    preferred = evidence.root / "results/report_figures" / filename
    return preferred if preferred.is_file() else evidence.comparison_path.parent / fallback


def build_pages(e: Evidence) -> list[tuple[str, list[Block]]]:
    a, b, c = e.rows
    labels = e.labels
    cfg = e.runs[0]["config"]
    precision_label, precision_description = training_precision(cfg)
    delta_b = finite(b["dice"]) - finite(a["dice"])
    delta_c = finite(c["dice"]) - finite(b["dice"])
    best = max(e.rows, key=lambda row: finite(row["dice"]))
    comparison_sentence = (
        f"Trên {e.comparison['image_count']} ảnh test, {best['experiment']} có mean Dice cao nhất "
        f"là {score(best['dice'])}. Chênh lệch Dice B so với A là {delta_b:+.4f}; "
        f"C so với B là {delta_c:+.4f}. Đây là kết quả của một split và một seed, "
        "chưa đủ để kết luận ưu thế ổn định qua nhiều lần chạy."
    )
    roles = [
        ["Nguyễn Bảo Duy", "24DH190309", "U-Net và tích hợp", "duy"],
        ["Trang Mai Quốc Khánh", "24DH190081", "Dữ liệu và kiểm tra dữ liệu", "khanh"],
        ["Nguyễn Hoàng Phúc", "24DH190314", "U-Net++ và sơ đồ kiến trúc", "phuc"],
        ["Lê Trần Nhật Huy", "24DH190312", "Loss và trainer dùng chung", "huy"],
        ["Nguyễn Đức Trí", "24DH190272", "Đánh giá và demo", "tri"],
    ]
    split_rows = []
    for split, title in (("train", "Train"), ("val", "Validation"), ("test", "Test")):
        split_rows.append([title, str(e.audit["split_counts"][split]),
                           *[str(e.audit["area_group_counts"][name][split]) for name in ("small", "medium", "large")]])
    stages = [
        ["Ảnh RGB", "3", "256 × 256", "Chuẩn hóa /255"],
        ["Encoder 0", "32", "256 × 256", "DoubleConv"],
        ["Encoder 1", "64", "128 × 128", "Pool và DoubleConv"],
        ["Encoder 2", "128", "64 × 64", "Pool và DoubleConv"],
        ["Encoder 3", "256", "32 × 32", "Pool và DoubleConv"],
        ["Bottleneck", "512", "16 × 16", "Pool và DoubleConv"],
        ["Decoder 3", "256", "32 × 32", "Up concat skip DoubleConv"],
        ["Decoder 2", "128", "64 × 64", "Up concat skip DoubleConv"],
        ["Decoder 1", "64", "128 × 128", "Up concat skip DoubleConv"],
        ["Decoder 0", "32", "256 × 256", "Up concat skip DoubleConv"],
        ["Head logits", "1", "256 × 256", "Conv 1 × 1"],
    ]
    architecture_dir = e.root / "members/phuc/works_done/artifacts/architecture"
    pages: list[tuple[str, list[Block]]] = []
    pages.append(("", [
        ("title", TITLE),
        p("Ngày tạo báo cáo " + author_date()),
        p("Báo cáo đồ án môn Deep Learning của nhóm năm thành viên. Nội dung gồm kiến trúc U-Net++, cơ chế nested dense skip connections, triển khai phân đoạn polyp trên Kvasir-SEG và so sánh với U-Net chuẩn."),
        p(comparison_sentence),
        h("Phạm vi đồ án"),
        p("A là U-Net baseline, B là U-Net++ dùng một head cuối. A/B trả lời yêu cầu cốt lõi của đề tài. C bổ sung deep supervision để khảo sát việc huấn luyện nhiều head trong cùng U-Net++. C suy luận bằng head cuối; đề tài này không thực hiện average-head accurate mode hay pruning trong paper."),
        table(["Thành viên", "MSSV", "Phần phụ trách", "Folder"], roles, [2.1, 1.15, 2.1, 1.15]),
        p("Các folder thể hiện phân công mô-đun trong một repo chung. Báo cáo và code hỗ trợ học, kiểm tra và trình bày; từng thành viên cần tự chạy, giải thích và xác nhận phần mình trước khi nộp."),
        h("Bằng chứng đi kèm"),
        p(f"Dữ liệu được audit; {e.tests['tests_run']} kiểm tra tự động đạt; ba run đã kết thúc; best checkpoint được khóa trước test. Bảng kết quả lấy trực tiếp từ JSON/CSV của evaluator, không dùng điểm từ bài thử học thuộc hai ảnh train."),
    ]))
    pages.append(("1 Bài toán và dữ liệu", [
        p("Input là một ảnh nội soi RGB. Output là mask nhị phân có cùng kích thước xử lý: 1 cho pixel polyp, 0 cho nền. Đây là semantic segmentation theo pixel; không dự đoán chẩn đoán, loại mô hay tọa độ bounding box."),
        p("Kvasir-SEG cung cấp ảnh polyp và mask phân đoạn tương ứng [3]. Trong repo, audit xác nhận 1.000 cặp ảnh/mask, không có mask rỗng, không có lỗi đọc hoặc lệch kích thước, và không phát hiện ảnh trùng hoàn toàn theo pixel đã giải mã. Tính độc lập theo bệnh nhân/video và ảnh gần trùng chưa được xác minh."),
        table(["Tập", "Số ảnh", "Small", "Medium", "Large"], split_rows, [1.9, 1.2, 1.2, 1.2, 1.2], (1, 2, 3, 4)),
        p("Phân tầng dùng tỷ lệ diện tích mask trên ảnh gốc: small dưới 5%, medium từ 5% đến dưới 20%, large từ 20% trở lên. Đây là nhóm diện tích trên ảnh, không phải kích thước polyp đo bằng milimét. Cột kfold_5 trong manifest không được dùng trong thực nghiệm holdout này."),
        h("Tiền xử lý dùng chung"),
        p("Đọc byte rồi giải mã OpenCV để hỗ trợ đường dẫn tiếng Việt trên Windows. Chuyển BGR thành RGB; ảnh resize linear về 256 × 256 và chia 255. Mask JPEG được nhị phân bằng grayscale >127 trước resize nearest; tensor mask có giá trị 0/1 và shape [B,1,256,256]."),
        p("Chỉ train có augmentation đồng bộ cho ảnh/mask: lật ngang, lật dọc, xoay 90 độ, affine scale 0,9–1,1, tịnh tiến ±6%, xoay ±15 độ. ColorJitter chỉ đổi màu ảnh. Validation/test chỉ resize và normalize. Seed riêng của Albumentations và sampler được reset theo seed cộng epoch để A/B/C gặp cùng thứ tự ảnh và chuỗi biến đổi."),
        p("Train cập nhật trọng số. Validation chọn epoch và dừng sớm. Test chỉ đo sau khi khóa cấu hình, threshold và checkpoint. Trainer không tạo test loader; test ID và SHA256 manifest giữ nguyên cho cả ba model."),
    ]))
    pages.append(("2 Kiến trúc UNet", [
        p("U-Net dùng contracting path để học ngữ cảnh và expanding path để định vị vùng phân đoạn [1]. Bản triển khai giữ topology encoder-decoder với bốn skip trực tiếp. Same padding, base 32 và ConvTranspose2d là lựa chọn đã ghi trong protocol; đây không phải tái lập nguyên xi mọi chi tiết của U-Net 2015."),
        table(["Vị trí", "Channel", "H và W", "Phép xử lý"], stages, [1.6, 0.8, 1.1, 3.2], (1, 2)),
        p("DoubleConv gồm hai Conv 3 × 3, padding=1 và ReLU. MaxPool2d giảm H/W một nửa. ConvTranspose2d kernel=2, stride=2 tăng resolution và chiếu channel; torch.cat dim=1 nối feature encoder với decoder. DoubleConv sau concat học cách phối hợp thông tin vị trí và ngữ cảnh."),
        p(f"Head Conv 1 × 1 đưa 32 channel thành một channel logits. Với cấu hình mặc định, U-Net có {integer(a['parameters'])} tham số. Model không có sigmoid trong forward: BCEWithLogitsLoss cần logits để tính loss ổn định; sigmoid được dùng khi tính soft Dice và suy luận [4]."),
        p("Điều kiện giao diện: input float NCHW, H/W là bội của 16 và ít nhất 16; output một tensor logits [B,1,H,W]. Kiểm tra shape và finite gradients phát hiện lỗi nối skip, nhầm channel hoặc normalize sai trước huấn luyện thật."),
    ]))
    pages.append(("3 Sơ đồ UNet", [
        figure(architecture_dir / "unet_word_portrait.png", "Hình 1. Topology của U-Net đang chạy trong repo. Kích thước là tensor của input 256 × 256 và base 32.", 5.2, 8.1),
    ]))
    pages.append(("4 Kiến trúc UNetPlusPlus", [
        p("U-Net++ thiết kế các đường skip lồng nhau và nối dày để giảm semantic gap giữa feature encoder và decoder [2]. Semantic gap chỉ khác biệt về mức biểu diễn của feature; cùng kích thước H/W không bảo đảm chúng đã biểu diễn thông tin tương tự. Đồ án kiểm chứng kết quả bằng mask, không đo trực tiếp semantic gap bằng một chỉ số riêng."),
        h("Các node trong code"),
        p("Ký hiệu X(i,j): i là mức độ sâu/resolution, j là cột tinh chỉnh. j=0 thuộc encoder. H là DoubleConv; Up là ConvTranspose2d từ mức i+1 lên mức i; ngoặc vuông là concat theo channel. Mỗi node có tham số riêng."),
        ("code", "X(i,0) = H(Pool(X(i-1,0)))  khi i > 0\nX(i,j) = H([X(i,0), ..., X(i,j-1),\n            Up(X(i+1,j-1))])  khi j > 0\nX(0,0) = H(input)"),
        p("Ví dụ X(0,1) nhận X(0,0) và Up(X(1,0)); X(0,2) nhận X(0,0), X(0,1) và Up(X(1,1)). Tại X(0,4), bốn node X(0,0)…X(0,3) được giữ lại trước khi nối với Up(X(1,3)). Đây là dense skip: node mới nhận toàn bộ node trước ở cùng level."),
        p("NestedSkipBlock chiếu feature bên dưới về C(i) channel. Vì có j node trước và một node upsample, tensor concat có (j+1) × C(i) channel. Ví dụ X(0,4) nhận 5 × 32 = 160 channel rồi DoubleConv đưa về 32. B lấy logits từ head của X(0,4)."),
        h("Deep supervision"),
        p("C thêm head tại X(0,1), X(0,2), X(0,3), X(0,4). Trainer tính loss cho cả bốn head rồi lấy trung bình, nên thang loss không tăng bốn lần. Validation/test chỉ dùng X(0,4). Trunk và head cuối B/C được khởi tạo giống nhau ở cùng seed; C chỉ thêm head phụ."),
        p(f"B có {integer(b['parameters'])} tham số; C có {integer(c['parameters'])}. A/B khác số tham số, nên kết quả này đồng thời phản ánh topology và dung lượng model. C không phải triển khai inference trung bình bốn head hoặc mạng đã pruning trong paper."),
    ]))
    pages.append(("5 Sơ đồ UNetPlusPlus", [
        figure(architecture_dir / "unetplusplus_word_portrait.png", "Hình 2. Các node nested dense skip. B dùng head cuối; C huấn luyện bốn head nhưng suy luận vẫn dùng head cuối.", 5.2, 8.1),
    ]))
    pages.append(("6 Loss và vòng huấn luyện", [
        p("Gọi z là logit, p = sigmoid(z), y là nhãn 0/1, N là số pixel của một ảnh. BCE đo lỗi từng pixel; soft Dice đo độ chồng lấp bằng xác suất để có gradient. Các biểu thức dưới đây được tính từng ảnh, rồi mean trong batch; epsilon = 10⁻⁶."),
        ("equation", "bce"),
        ("equation", "soft_dice"),
        ("equation", "combined"),
        ("equation", "deep_supervision"),
        p("Trong code, BCE được tính bằng binary_cross_entropy_with_logits, không đưa p qua sigmoid rồi gọi BCEWithLogits lần nữa. Thành phần soft Dice dùng p. Loss kiểm tra mask chỉ có 0/1, shape logits khớp mask, và mọi giá trị hữu hạn."),
        h("Một batch train"),
        p("model.train → lấy batch và chuyển lên GPU → optimizer.zero_grad → forward → loss → backward → optimizer.step. Autograd đi qua các phép toán để tính gradient theo trọng số; Adam dùng gradient cập nhật tham số. " + precision_description + " Loss nội bộ chuyển về float32 để phép tổng của Dice ổn định."),
        p("Sau mỗi epoch, model.eval và inference_mode đọc validation bằng FP32. Validation loss là cùng objective của head cuối cho A/B/C. Chọn best bằng validation Dice cao nhất; nếu hòa thì lấy validation loss thấp hơn. Dừng khi 10 epoch liên tiếp không cải thiện hoặc đạt 50 epoch."),
        p("best.pt lưu epoch được chọn; last.pt lưu epoch mới nhất để resume. Checkpoint chứa model_state, optimizer_state, scaler_state, config, lịch sử, RNG và provenance. Lưu tạm rồi thay file để giảm nguy cơ checkpoint hỏng khi tiến trình bị ngắt [6]."),
    ]))
    pages.append(("7 Giao thức so sánh", [
        table(["Thiết lập", "Giá trị dùng chung"], [
            ["Dữ liệu", "Manifest 700 train /150 validation /150 test; cùng ID"],
            ["Input", "RGB 256 × 256 /255; mask nhị phân resize nearest"],
            ["Model", "Base 32; 32/64/128/256/512; same padding; không BatchNorm"],
            ["Khởi tạo", "PyTorch mặc định; không pretrained, attention hoặc residual"],
            ["Optimizer", "Adam; learning rate 0,001; không scheduler"],
            ["Train", "Batch 8; tối đa 50 epoch; patience 10; seed 42; num_workers=0"],
            ["Loss", "0,5 BCEWithLogits + 0,5 soft Dice; epsilon=10⁻⁶"],
            ["Precision", f"{precision_label} trong train; FP32 validation/test; không TF32"],
            ["Đánh giá", "Sigmoid ≥0,5; head cuối; metric từng ảnh rồi mean"],
            ["Thay đổi A/B/C", "Kiến trúc U-Net so với U-Net++; thêm supervision ở C"],
        ], [1.65, 5.05]),
        p(f"Thiết bị thực nghiệm: {e.tests.get('gpu')}; Python {e.tests['python']}; PyTorch {e.tests['torch']}. Seed và deterministic algorithms hỗ trợ lặp lại trong môi trường đã ghi, nhưng không bảo đảm bitwise giống trên mọi phiên bản, thiết bị hay nền tảng [7]."),
        *([p("Lịch sử chỉnh precision: một run thử của C gặp overflow ở forward FP16 tại batch 28 của epoch 30. Kiểm tra trên dữ liệu train cho thấy trọng số checkpoint và forward FP32 vẫn hữu hạn. Nhóm giữ log thử, đổi precision dùng chung sang BF16 và chạy lại cả A/B/C từ đầu trước khi khóa checkpoint và đọc test. Đây là sửa lỗi ổn định số bằng chẩn đoán train, không điều chỉnh mô hình theo điểm test.")] if cfg.get("amp") and cfg.get("amp_dtype") == "bfloat16" else []),
        h("Khóa trước test"),
        p("configs/protocol_lock.json ghi checkpoint best, SHA256 checkpoint, config, manifest và thời điểm khóa. Evaluator đối chiếu cấu hình cùng mã preprocessing/model trước khi đọc test. File lock tạo dấu vết kiểm toán trong repo; tự sửa file hoặc timestamp không chứng minh lịch sử khách quan bên ngoài repo."),
        p("SHA256 manifest: " + e.comparison["manifest_sha256"]),
        p("SHA256 protocol lock: " + e.comparison["protocol_lock_sha256"]),
        p("Checkpoint được chọn bằng validation, threshold cố định trước test. Không chọn lại epoch, augmentation hay threshold sau khi xem điểm test. Nếu thiết kế thí nghiệm khác, phải tạo run và output mới để giữ bằng chứng cũ."),
    ]))
    pages.append(("8 Chỉ số và kết quả test", [
        p("Gọi TP là pixel polyp dự đoán đúng, FP là nền bị dự đoán thành polyp, FN là polyp bị bỏ sót. Dice và IoU đo chồng lấp vùng. Precision = TP/(TP+FP); recall = TP/(TP+FN). Vì nền có thể chiếm phần lớn ảnh, pixel accuracy không phải chỉ số kết luận chính."),
        ("equation", "dice"), ("equation", "iou"),
        p("Chỉ số được tính trên từng ảnh 256 × 256 rồi lấy trung bình 150 ảnh. Không gộp toàn bộ pixel thành một ảnh lớn. Dice và IoU của một ảnh có quan hệ Dice = 2IoU/(1+IoU), nhưng không thay mean IoU vào công thức để suy ra mean Dice."),
        table(["Model", "Dice", "IoU", "Precision", "Recall", "Epoch"],
              [[row["experiment"], *[score(row[name]) for name in ("dice", "iou", "precision", "recall")], str(row["selected_epoch"])] for row in e.rows],
              [1.3, 1.05, 1.05, 1.05, 1.05, 1.2], (1, 2, 3, 4, 5)),
        p(comparison_sentence),
        figure(compact_figure(e, "metric_comparison.png", "comparison.png"), "Hình 3. Mean Dice và IoU trên cùng 150 ảnh test. Giá trị lấy từ comparison.csv.", 6.35),
        p("Cả prediction và mask rỗng được quy ước Dice/IoU bằng 1. Nếu chỉ một vùng rỗng thì Dice/IoU bằng 0. Audit hiện không có mask gốc rỗng; trường hợp rỗng vẫn có test để kiểm tra logic metric."),
    ]))
    training_rows = []
    for label, run, history in zip(labels, e.runs, e.histories):
        peaks = [finite(row["peak_gpu_memory_bytes"]) for row in history if row.get("peak_gpu_memory_bytes")]
        training_rows.append([label, str(run["epochs_completed"]), str(run["best"]["epoch"]),
                              score(run["best"]["val_dice"]), f"{run['training_and_validation_seconds']/60:.2f}",
                              f"{max(peaks)/(1024**3):.2f}" if peaks else "Không ghi"])
    b_epoch10 = next((row for row in e.histories[1] if int(row["epoch"]) == 10), None)
    pause_note = (
        f" Cụ thể, epoch 10 của {labels[1]} ghi {finite(b_epoch10['epoch_seconds']) / 3600:.2f} giờ wall "
        "vì máy ngủ rồi tiến trình tiếp tục. Giá trị này bao gồm khoảng nghỉ của máy."
        if b_epoch10 and finite(b_epoch10["epoch_seconds"]) > 3600 else ""
    )
    pages.append(("9 Diễn biến huấn luyện và chi phí", [
        table(["Model", "Epoch đã chạy", "Epoch best", "Val Dice best", "Phút wall qua epoch", "VRAM đỉnh GiB"], training_rows,
              [0.9, 1.0, 1.0, 1.25, 1.35, 1.2], (1, 2, 3, 4, 5)),
        p("Epoch best có thể sớm hơn epoch cuối. Các run dừng theo cùng patience và giới hạn tối đa; số epoch thực tế có thể khác nhau. Đường validation dùng chọn checkpoint, chưa phải đường test và không được gọi là khả năng tổng quát hóa trên dữ liệu ngoài Kvasir-SEG."),
        p("Trường training_and_validation_seconds là tổng wall time của các epoch đã lưu, không phải thời gian GPU tính toán. Nó có thể gồm thời gian máy hoặc tiến trình bị tạm dừng trong một epoch; khoảng nghỉ giữa hai lần resume ngoài epoch không được cộng vào tổng này." + pause_note + " Không dùng bảng phút wall để kết luận model nào train nhanh hơn. So sánh tốc độ riêng bằng phép đo forward trên cùng GPU và FP32 ở bảng dưới."),
        table(["Model", "Tham số", "Forward ms mỗi ảnh", "Forward ảnh mỗi giây"],
              [[row["experiment"], integer(row["parameters"]), f"{finite(row['forward_ms_per_image']):.3f}", f"{finite(row['forward_images_per_second']):.1f}"] for row in e.rows],
              [1.0, 1.6, 2.05, 2.05], (1, 2, 3)),
        p("Timing chỉ đo model forward sau warmup, cùng phần cứng, batch và FP32. Không gồm đọc/resize ảnh, truyền CPU→GPU, sigmoid, metric hay xuất mask. Do đó đây không phải FPS của demo đầu cuối. VRAM đỉnh là bộ nhớ torch đã cấp phát trong train, không phải toàn bộ VRAM của máy."),
        p("C tính cả bốn head trong forward theo giao diện hiện tại; evaluator chọn head cuối sau forward. Chi phí forward C vẫn gồm head phụ, không coi đây là chi phí một model đã pruning. Timing chỉ đo một lượt và có thể dao động theo trạng thái GPU. Dù số đo C trong lượt này thấp hơn B, hai model có cùng graph chính và C tính thêm ba head phụ; chưa thể kết luận deep supervision giảm thời gian suy luận."),
    ]))
    pages.append(("10 Đường học của ba model", [
        p("Đường học giúp đối chiếu loss trong train với validation Dice dùng chọn best. Các epoch sau best vẫn được lưu để xác định dừng sớm; báo cáo test dùng best checkpoint, không tự động dùng last checkpoint."),
        figure(compact_figure(e, "training_curves.png", "training_curves.png"), "Hình 4. Loss train và validation cùng validation Dice theo epoch. Train loss C là mean bốn head; val loss C chỉ dùng head cuối.", 6.45, 7.0),
        p("Với C, train loss và val loss khác số head được giám sát, đồng thời train có augmentation còn validation không có. Vì vậy không diễn giải khoảng cách giữa hai đường chỉ bằng overfitting. Cần đọc thêm validation Dice, best epoch và mask trước khi đưa ra giả thuyết về lỗi học."),
    ]))
    group_rows = []
    for category in ("small", "medium", "large"):
        selected = [next(row for row in e.groups if row["experiment"] == label and row["polyp_category"] == category) for label in labels]
        group_rows.append([category, selected[0]["count"], *[score(row["dice"]) for row in selected], *[score(row["iou"]) for row in selected]])
    paired_rows = []
    for label in labels[1:]:
        diff = e.comparison["paired_differences"][label]
        gained, worse = int(diff["improved_images"]), int(diff["worse_images"])
        paired_rows.append([f"{label} so với {labels[0]}", f"{diff['mean_delta_dice']:+.4f}", str(gained), str(worse), str(len(e.paired) - gained - worse)])
    selection_text = (
        "Hình 5 của bản Word dùng ba test ID đầu theo thứ tự sắp xếp, không chọn theo chất lượng mask. "
        "File paired_examples.png trong comparison lưu phân tích cực trị riêng: hai delta B−A thấp nhất, "
        "hai cao nhất và hai mean Dice thấp nhất giữa ba model, rồi bỏ ID trùng."
        if (e.root / "results/report_figures/paired_examples.png").is_file() else
        "Hình 5 chọn sau đánh giá: hai delta B−A thấp nhất, hai delta B−A cao nhất và hai ảnh có mean Dice "
        "thấp nhất giữa ba model, sau đó bỏ ID trùng."
    )
    pages.append(("11 Phân tích theo ảnh và diện tích", [
        table(["Nhóm", "n", "Dice A", "Dice B", "Dice C", "IoU A", "IoU B", "IoU C"], group_rows,
              [0.85, 0.45, 0.9, 0.9, 0.9, 0.9, 0.9, 0.9], (1, 2, 3, 4, 5, 6, 7)),
        p("Nhóm được gán bằng diện tích mask trên ảnh gốc. Metric vẫn tính ở resolution 256 × 256. Small có ít pixel dương; một lỗi vài pixel ở biên có thể ảnh hưởng Dice mạnh hơn. Phân tầng giúp đọc các nhóm khó, nhưng không đo kích thước lâm sàng hoặc thay thế đánh giá theo bệnh nhân."),
        table(["So sánh paired", "Mean delta Dice", "Tăng", "Giảm", "Bằng"], paired_rows,
              [2.25, 1.45, 1.0, 1.0, 1.0], (1, 2, 3, 4)),
        p("Paired analysis ghép đúng cùng image_id rồi trừ Dice của model tham chiếu. Điểm trung bình có thể tăng trong khi vẫn có nhiều ảnh giảm. Số ảnh tăng/giảm mô tả split này; không phải kiểm định ý nghĩa thống kê và không có độ lệch chuẩn qua nhiều seed."),
        h("Đọc các ca khó"),
        p(selection_text + " Các ví dụ dùng phân tích lỗi; không được dùng để sửa cấu hình rồi báo lại test như một đánh giá độc lập."),
        p("Khi đọc mask, kiểm tra vùng polyp bị bỏ sót, phần nền bị chọn nhầm, sai biên và mảnh nhỏ rời vùng chính. Có thể thảo luận ảnh phản xạ mạnh, biên mờ hoặc tương phản thấp như giả thuyết nguyên nhân, nhưng không gán nhãn nguyên nhân chắc chắn chỉ từ Dice."),
        h("Giới hạn thực nghiệm"),
        p("Một seed 42; một holdout 700/150/150; số tham số A/B khác nhau; chưa xác minh near duplicates và patient/video independence; mask JPEG và resize có thể đổi chi tiết biên. Điểm trên Kvasir-SEG này không chứng minh độ tin cậy lâm sàng hoặc khả năng chạy tốt ở cơ sở y tế khác."),
    ]))
    pages.append(("12 Minh họa dự đoán và lỗi", [
        figure(compact_figure(e, "paired_examples.png", "paired_examples.png"), "Hình 5. Ảnh, ground truth và prediction của A/B/C trên các test ID được chọn theo quy tắc đã ghi. Dice hiển thị từng ảnh.", 6.65),
        p("Đọc từng hàng theo cùng ảnh và cùng GT. Vùng dự đoán gần đúng về tổng diện tích vẫn có thể sai vị trí, nên phải xem mask chồng lấp. Nếu model bỏ sót polyp, recall giảm; nếu chọn cả nền, precision giảm. Dice tóm tắt hai lỗi đó nhưng không chỉ rõ sai tại pixel nào."),
        p("Toàn bộ kết quả từng ảnh được lưu trong paired_per_image.csv và per_image.csv của từng evaluator. Các mask test giữ dạng PNG 0/255 ở resolution input. Demo resize bản đồ xác suất về ảnh gốc bằng bilinear rồi threshold 0,5 để tạo mask hiển thị; điểm trong báo cáo vẫn đo ở input 256 × 256."),
        p("Một hình đẹp không đại diện toàn bộ 150 ảnh. Cần xem cả bảng paired và các ví dụ cực trị trong output comparison để đọc biến thiên và thất bại. Kết luận chính dựa trên toàn bộ test, cùng những giới hạn ở mục 11."),
    ]))
    command_text = "\n".join([
        "$PolypPython = '..\\.venv-polyp\\Scripts\\python.exe'",
        "& $PolypPython -X utf8 -B `",
        "  -m members.duy.works_done.verify_project",
        "& $PolypPython -X utf8 -B `",
        "  -m members.khanh.works_done.audit_data",
        "& $PolypPython -X utf8 -B -m members.huy.works_done.train `",
        "  --config configs/unet.json --output runs/new_A_seed42",
        "# Lặp lại với unetplusplus.json và unetplusplus_ds.json.",
        "# Khóa cả ba best checkpoint trước khi đọc test.",
        "& $PolypPython -X utf8 -B `",
        "  -m members.duy.works_done.lock_protocol `",
        "  --runs runs/new_A_seed42 runs/new_B_seed42 runs/new_C_seed42 `",
        "  --labels A B C --output configs/new_protocol_lock.json",
        "& $PolypPython -X utf8 -B `",
        "  -m members.tri.works_done.evaluate `",
        "  --checkpoint runs/new_A_seed42/best.pt --split test `",
        "  --protocol-lock configs/new_protocol_lock.json `",
        "  --output results/new_test_A --device cuda",
        "# Lặp evaluator cho B/C; compare ba output mới.",
        "& $PolypPython -X utf8 -B `",
        "  -m members.tri.works_done.compare `",
        "  --evaluations results/new_test_A results/new_test_B `",
        "                results/new_test_C --output results/new_comparison",
    ])
    pages.append(("13 Tái lập và tổ chức code", [
        p("Chạy từ repo root sau khi tạo môi trường và cài đúng requirements khóa. Các output bên dưới dùng tên mới để giữ thí nghiệm đã có. $PolypPython trỏ đến môi trường thực nghiệm hiện tại; trên máy mới có thể đổi sang .venv/Scripts/python.exe. Report generator dùng runtime tài liệu riêng và không cần torch."),
        ("code", command_text),
        h("Truy ngược kết quả"),
        p("results/data_audit/audit.json xác nhận dữ liệu; results/verification/test_summary.json ghi kiểm tra; runs/*/history.csv ghi diễn biến; best.pt chứa checkpoint chọn bằng validation; configs/protocol_lock.json khóa lựa chọn; results/comparison/ chứa bảng và hình cuối."),
        p("Provenance train ghi commit, branch, trạng thái dirty, SHA256 từng file chạy và phiên bản runtime. Commit ở thời điểm train: " + e.runs[0]["provenance"]["source_commit"] + ". Khi dirty=true, chỉ commit không đủ tái lập; phải giữ các source hashes và đúng code trong gói bàn giao."),
        p("Mỗi người code ở members/<tên>/works_done. Duy quản lý model factory và tích hợp; các mô-đun đọc dữ liệu, trainer, metrics và checkpoint IO có giao diện chung. Không sao chép Dataset hoặc công thức metric thành nhiều phiên bản khác nhau."),
    ]))
    pages.append(("14 Kết luận và chuẩn bị bảo vệ", [
        p(comparison_sentence),
        p("Đồ án đã triển khai dữ liệu, hai kiến trúc, trainer chung, deep supervision mở rộng, đánh giá khóa trước test và demo. Kết quả cho phép đối chiếu độ chồng lấp, số tham số và chi phí forward trong điều kiện đã ghi. Bước nghiên cứu tiếp theo hợp lý là lặp nhiều seed, kiểm tra phân chia theo bệnh nhân/video và đánh giá trên dataset ngoài; các bước này chưa nằm trong kết quả hiện tại."),
        h("Trình bày trong khoảng năm phút"),
        p("Phút 1 nêu input, mask và split. Phút 2 vẽ U-Net và giải thích skip. Phút 3 chỉ X(0,2) và X(0,4) trong U-Net++, rồi phân biệt nested dense skip với deep supervision. Phút 4 trình bày loss, checkpoint và bảng test. Phút 5 chỉ một ca thất bại và giới hạn một seed."),
        h("Những câu phải tự trả lời được"),
        p("Vì sao output là một channel logits? Vì sao mask resize nearest? X(0,4) có bao nhiêu channel sau concat? Vì sao C mean bốn loss? Validation khác test ở bước nào? Mean Dice khác metric gộp pixel thế nào? Vì sao B nhiều tham số hơn A? Số liệu nào chứng minh model đã train thay vì chỉ chạy forward? Vì sao bài thử học thuộc hai ảnh không phải kết quả cuối?"),
        p("Duy cần tự chạy một lần predict, mở history, truy ngược best epoch và tính Dice trên ví dụ nhỏ. Mỗi bạn cần giải thích được tensor đi qua folder mình phụ trách. Khả năng trình bày phải được xác nhận bằng việc tự thực hành; code hoặc báo cáo được hỗ trợ tạo không tự chứng minh nhóm đã hiểu."),
        h("Kiểm tra trước khi nộp"),
        p("Đối chiếu yêu cầu giảng viên với A/B; kiểm tra tên thành viên; giữ checkpoint, manifest, config, CSV và source hashes; chạy demo bằng đúng best checkpoint; báo điểm test cùng split/seed; không biến giả thuyết lỗi thành kết luận y khoa; ghi rõ C dùng final head."),
    ]))
    pages.append(("15 Tài liệu tham khảo", [
        *[("reference", reference) for reference in REFERENCES],
        h("Bằng chứng thực nghiệm trong repo"),
        p("Số liệu dữ liệu, kiểm tra, checkpoint, metric và timing trong báo cáo lấy từ artifact của repo, không lấy từ điểm của paper. Paper cung cấp cơ sở kiến trúc và dataset; chi tiết same padding, base 32, loss weights, seed, AMP và head policy thuộc protocol triển khai này."),
        p("Bảng chính: results/comparison/comparison.csv. Nhóm diện tích: comparison_by_size.csv. Kết quả theo ảnh: paired_per_image.csv. Các file JSON giữ số chưa làm tròn; bảng báo cáo làm tròn metric đến bốn chữ số thập phân."),
        p("SHA256 comparison JSON: " + sha256(e.comparison_path)),
    ]))
    return pages


def omml_run(text: str):
    """Native Word math run, keeping formulas editable and structured."""
    from docx.oxml import OxmlElement
    run = OxmlElement("m:r")
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    return run


def omml_fraction(numerator: list, denominator: list):
    from docx.oxml import OxmlElement
    fraction = OxmlElement("m:f")
    for name, elements in (("m:num", numerator), ("m:den", denominator)):
        container = OxmlElement(name)
        for element in elements:
            container.append(element)
        fraction.append(container)
    return fraction


def omml_subscript(base: str, subscript: str):
    from docx.oxml import OxmlElement
    result = OxmlElement("m:sSub")
    for tag, text in (("m:e", base), ("m:sub", subscript)):
        container = OxmlElement(tag)
        container.append(omml_run(text))
        result.append(container)
    return result


def omml_sum(expression: str, index: str = "j=1", upper: str = "N"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    nary = OxmlElement("m:nary")
    properties = OxmlElement("m:naryPr")
    symbol = OxmlElement("m:chr")
    symbol.set(qn("m:val"), "∑")
    properties.append(symbol)
    nary.append(properties)
    for tag, text in (("m:sub", index), ("m:sup", upper), ("m:e", expression)):
        container = OxmlElement(tag)
        container.append(omml_run(text))
        nary.append(container)
    return nary


def equation_elements(kind: str) -> list:
    r = omml_run
    sub = omml_subscript
    if kind == "bce":
        return [sub("L", "BCE"), r(" = −"), omml_fraction([r("1")], [r("N")]),
                omml_sum("[yⱼ log(pⱼ) + (1−yⱼ) log(1−pⱼ)]")]
    if kind == "soft_dice":
        return [sub("D", "soft"), r(" = "), omml_fraction(
            [r("2 "), omml_sum("pⱼ yⱼ"), r(" + ε")],
            [omml_sum("pⱼ"), r(" + "), omml_sum("yⱼ"), r(" + ε")])]
    if kind == "combined":
        return [r("L = 0,5 "), sub("L", "BCE"), r(" + 0,5 (1 − "), sub("D", "soft"), r(")")]
    if kind == "deep_supervision":
        return [sub("L", "DS"), r(" = "), omml_fraction([r("1")], [r("4")]), omml_sum("Lₖ", "k=1", "4")]
    if kind == "dice":
        return [r("Dice = "), omml_fraction([r("2 TP")], [r("2 TP + FP + FN")])]
    if kind == "iou":
        return [r("IoU = "), omml_fraction([r("TP")], [r("TP + FP + FN")])]
    raise ValueError(f"Unknown equation: {kind}")


EQUATIONS_TEXT = {
    "bce": "L_BCE = −(1/N) Σⱼ [yⱼ log(pⱼ) + (1−yⱼ) log(1−pⱼ)]",
    "soft_dice": "D_soft = (2 Σⱼ pⱼyⱼ + ε) / (Σⱼ pⱼ + Σⱼ yⱼ + ε)",
    "combined": "L = 0,5 L_BCE + 0,5 (1 − D_soft)",
    "deep_supervision": "L_DS = (1/4) Σₖ₌₁⁴ Lₖ",
    "dice": "Dice = 2TP / (2TP + FP + FN)",
    "iou": "IoU = TP / (TP + FP + FN)",
}


def set_black_font(run, size: float | None = None, name: str = "Calibri") -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    run.font.name = name
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().rFonts
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    properties.append(color)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def add_word_table(document, data: dict) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    word_table = document.add_table(rows=1, cols=len(data["headers"]))
    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    word_table.autofit = False
    for column, width in zip(word_table.columns, data["widths"]):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            element.set(qn(f"w:{key}"), value)
        borders.append(element)
    word_table._tbl.tblPr.append(borders)
    repeat = OxmlElement("w:tblHeader")
    word_table.rows[0]._tr.get_or_add_trPr().append(repeat)
    all_rows = [data["headers"], *data["rows"]]
    for row_index, values in enumerate(all_rows):
        cells = word_table.rows[0].cells if row_index == 0 else word_table.add_row().cells
        for column_index, (cell, value, width) in enumerate(zip(cells, values, data["widths"])):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), "90")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index in data["numeric"] else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            set_black_font(run, 10.0)
            if row_index == 0:
                run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EDEDED")
                properties.append(shading)
        prevent_split = OxmlElement("w:cantSplit")
        word_table.rows[row_index]._tr.get_or_add_trPr().append(prevent_split)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.1
    set_black_font(spacer.add_run(""), 2)


def png_size(path: Path) -> tuple[int, int]:
    data = path.read_bytes()[:24]
    if data[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"Report figures must be PNG: {path}")
    return struct.unpack(">II", data[16:24])


def write_word(pages: list[tuple[str, list[Block]]], output: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.3)
    for style_name, size in (("Normal", 11), ("Title", 20), ("Subtitle", 11), ("Heading 1", 14), ("Heading 2", 12), ("Caption", 10)):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), "Calibri")
        paragraph_properties = style.element.get_or_add_pPr()
        for border in list(paragraph_properties.findall(qn("w:pBdr"))):
            paragraph_properties.remove(border)
    normal = document.styles["Normal"].paragraph_format
    normal.space_after = Pt(7)
    normal.line_spacing = 1.08
    normal.widow_control = True
    for heading in ("Heading 1", "Heading 2"):
        fmt = document.styles[heading].paragraph_format
        fmt.space_before, fmt.space_after = Pt(8), Pt(7)
        fmt.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_black_font(footer.add_run("Trang "), 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = TITLE
    document.core_properties.subject = "Đồ án Deep Learning phân đoạn polyp trên Kvasir SEG"
    document.core_properties.author = "Nhóm đồ án Deep Learning"
    document.core_properties.keywords = "U-Net, U-Net++, Kvasir-SEG, polyp segmentation"
    for page_index, (heading, blocks) in enumerate(pages):
        if page_index:
            document.add_page_break()
        if heading:
            document.add_heading(heading, level=1)
        for kind, data in blocks:
            if kind in {"paragraph", "title", "subheading"}:
                if kind == "title":
                    paragraph = document.add_paragraph(data, style="Title")
                    paragraph.paragraph_format.space_after = Pt(14)
                elif kind == "subheading":
                    paragraph = document.add_heading(data, level=2)
                else:
                    paragraph = document.add_paragraph(data)
                for run in paragraph.runs:
                    set_black_font(run)
            elif kind == "table":
                add_word_table(document, data)
            elif kind == "code":
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(8)
                paragraph.paragraph_format.line_spacing = 1.0
                for index, line in enumerate(data.splitlines()):
                    if index:
                        paragraph.add_run().add_break()
                    set_black_font(paragraph.add_run(line), 9, "Consolas")
            elif kind == "equation":
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(7)
                native = OxmlElement("m:oMathPara")
                equation = OxmlElement("m:oMath")
                for element in equation_elements(data):
                    equation.append(element)
                native.append(equation)
                paragraph._p.append(native)
            elif kind == "figure":
                width, height = png_size(data["path"])
                desired_width = data["width"]
                # Fit full-page architecture and tall post-hoc figure without distortion.
                maximum_height = data["max_height"]
                effective_width = min(desired_width, maximum_height * width / height)
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                picture = paragraph.add_run().add_picture(str(data["path"]), width=Inches(effective_width))
                picture._inline.docPr.set("descr", data["caption"])
                caption = document.add_paragraph(data["caption"], style="Caption")
                caption.paragraph_format.space_after = Pt(7)
                for run in caption.runs:
                    set_black_font(run, 10)
            elif kind == "reference":
                number, label, url = data
                paragraph = document.add_paragraph(f"[{number}] {label}")
                paragraph.paragraph_format.keep_with_next = True
                link = document.add_paragraph()
                add_hyperlink(link, url, url)
                link.paragraph_format.space_after = Pt(8)
            else:
                raise ValueError(f"Unknown block: {kind}")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def write_markdown(pages: list[tuple[str, list[Block]]], output: Path, root: Path) -> None:
    lines: list[str] = []
    for heading, blocks in pages:
        if heading:
            lines.extend([f"## {heading}", ""])
        for kind, data in blocks:
            if kind == "title":
                lines.extend([f"# {data}", ""])
            elif kind == "subheading":
                lines.extend([f"### {data}", ""])
            elif kind == "paragraph":
                lines.extend([data, ""])
            elif kind == "table":
                def row_text(row):
                    return "| " + " | ".join(str(value).replace("|", "\\|") for value in row) + " |"
                lines.append(row_text(data["headers"]))
                lines.append(row_text(["---"] * len(data["headers"])))
                lines.extend(row_text(row) for row in data["rows"])
                lines.append("")
            elif kind == "code":
                lines.extend(["```text", data, "```", ""])
            elif kind == "equation":
                lines.extend([f"`{EQUATIONS_TEXT[data]}`", ""])
            elif kind == "figure":
                # Markdown lives in docs/, so figures point one level up.
                relative = "../" + data["path"].relative_to(root).as_posix()
                lines.extend([f"![{data['caption']}]({relative})", ""])
            elif kind == "reference":
                number, label, url = data
                lines.extend([f"[{number}] [{label}]({url})", ""])
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--output", type=Path, default=Path("reports"))
    parser.add_argument("--comparison", type=Path, default=Path("results/comparison/comparison.json"))
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    root = args.project_root.resolve()
    comparison_path = args.comparison if args.comparison.is_absolute() else root / args.comparison
    evidence = load_evidence(root, comparison_path)
    pages = build_pages(evidence)
    if args.validate_only:
        print(json.dumps({"status": "evidence_validated", "planned_sections": len(pages),
                          "test_images": len(evidence.paired), "labels": evidence.labels}, ensure_ascii=False, indent=2))
        return
    output_dir = args.output if args.output.is_absolute() else root / args.output
    word = output_dir / "BAO_CAO_UNET_UNETPLUSPLUS_KVASIR_SEG.docx"
    markdown = root / "docs/BAO_CAO_DO_AN.md"
    write_word(pages, word)
    write_markdown(pages, markdown, root)
    print(json.dumps({"status": "report_created_pending_visual_qa", "docx": str(word),
                      "markdown": str(markdown), "planned_sections": len(pages)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
